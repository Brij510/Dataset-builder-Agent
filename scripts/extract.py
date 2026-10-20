#!/usr/bin/env python3
"""Stage 1 - extract clean text from raw documents into a cached corpus.

Process-based (pdfplumber is GIL-bound, threads give no speedup), resumable, and
caches per file so re-runs only parse new documents.

Usage:
  python extract.py --input ./docs --output ./dataset [--pages 4] [--workers 4]
"""
import argparse
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import clean_text, load_config, read_json, write_json  # noqa: E402


def extract_pdf(path, pages):
    import pdfplumber
    out = []
    try:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                if i >= pages:
                    break
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                if t.strip():
                    out.append(t)
    except Exception:
        return ""
    return "\n".join(out)


def extract_docx(path, pages):
    try:
        import docx
    except ImportError:
        return ""
    try:
        d = docx.Document(path)
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        return "\n".join(paras[:pages * 120])
    except Exception:
        return ""


def extract_plain(path, pages):
    try:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""


def extract_one(job):
    rel, path, pages, exts = job
    if exts == ".pdf":
        raw = extract_pdf(path, pages)
    elif exts == ".docx":
        raw = extract_docx(path, pages)
    else:
        raw = extract_plain(path, pages)
    return rel, clean_text(raw)


def collect_files(root, extensions):
    files = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for name in sorted(filenames):
            ext = os.path.splitext(name)[1].lower()
            if ext in extensions:
                files.append(os.path.join(dirpath, name))
    return sorted(files)


def main():
    ap = argparse.ArgumentParser(description="Extract text from raw documents")
    ap.add_argument("--input", required=True, help="root folder of source documents")
    ap.add_argument("--output", required=True, help="dataset working directory")
    ap.add_argument("--pages", type=int, default=4, help="pages to read per PDF")
    ap.add_argument("--workers", type=int, default=max(2, (os.cpu_count() or 4) - 1))
    ap.add_argument("--ext", default=".pdf,.docx,.txt,.md")
    ap.add_argument("--min-chars", type=int, default=300)
    ap.add_argument("--config", help="optional config file (values act as defaults)")
    args = ap.parse_args()

    cfg = load_config(args.config)
    if args.pages == 4:
        args.pages = cfg["extract"].get("pages", args.pages)
    if args.min_chars == 300:
        args.min_chars = cfg["extract"].get("min_chars", args.min_chars)

    exts = tuple(e.strip().lower() for e in args.ext.split(","))
    os.makedirs(args.output, exist_ok=True)
    corpus_path = os.path.join(args.output, "corpus.json")

    corpus = read_json(corpus_path, {}) or {}
    files = collect_files(args.input, exts)
    todo = [f for f in files
            if os.path.relpath(f, args.input) not in corpus]
    print(f"[extract] {len(files)} documents found, {len(todo)} to parse, "
          f"{args.workers} workers, {args.pages}p", flush=True)
    if not todo:
        print(f"[extract] corpus already complete -> {corpus_path}")
        return

    t0 = time.time()
    done = 0
    with ProcessPoolExecutor(max_workers=args.workers) as ex:
        jobs = [(os.path.relpath(f, args.input), f, args.pages, os.path.splitext(f)[1].lower())
                for f in todo]
        futs = [ex.submit(extract_one, j) for j in jobs]
        for fut in as_completed(futs):
            rel, text = fut.result()
            corpus[rel] = text
            done += 1
            if done % 25 == 0:
                write_json(corpus_path, corpus)
                rate = done / max(time.time() - t0, 0.001)
                print(f"  {done}/{len(todo)}  {rate:.1f} doc/s", flush=True)
    write_json(corpus_path, corpus)

    kept = {k: v for k, v in corpus.items() if len(v) >= args.min_chars}
    total = sum(len(v) for v in kept.values())
    print(f"[extract] done: {len(corpus)} parsed, {len(kept)} usable, "
          f"{total:,} chars, {time.time() - t0:.0f}s")
    print(f"[extract] corpus -> {corpus_path}")


if __name__ == "__main__":
    main()
