#!/usr/bin/env python3
"""Tests for document formats and configuration-driven behaviour."""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.abspath(os.path.join(HERE, "..", "scripts"))
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, HERE)

from fixtures import write  # noqa: E402

PY = sys.executable


def run(script, *args, expect_ok=True):
    cmd = [PY, os.path.join(SCRIPTS, script)] + [str(a) for a in args]
    p = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    if expect_ok and p.returncode != 0:
        raise AssertionError(f"{script} failed\n{p.stdout[-2000:]}\n{p.stderr[-2000:]}")
    return p


class DocxInput(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="dsb_docx_")
        self.src = os.path.join(self.tmp, "docs")
        self.out = os.path.join(self.tmp, "out")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_docx_paragraphs_are_extracted(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        os.makedirs(os.path.join(self.src, "Math"), exist_ok=True)
        path = os.path.join(self.src, "Math", "notes.docx")
        d = docx.Document()
        d.add_paragraph("A polynomial in one variable is an expression of a certain form.")
        d.add_paragraph("The degree of a polynomial is its highest power.")
        d.save(path)
        run("extract.py", "--input", self.src, "--output", self.out, "--workers", "1")
        corpus = json.load(open(os.path.join(self.out, "corpus.json")))
        self.assertTrue(corpus, "docx produced no text")
        joined = " ".join(corpus.values())
        self.assertIn("polynomial", joined)
        self.assertIn("degree", joined)

    def test_docx_tables_are_extracted(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        os.makedirs(os.path.join(self.src, "Science"), exist_ok=True)
        path = os.path.join(self.src, "Science", "table.docx")
        d = docx.Document()
        d.add_paragraph("The following table lists important acids.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Acid"
        t.cell(0, 1).text = "Formula"
        t.cell(1, 0).text = "Hydrochloric acid"
        t.cell(1, 1).text = "HCl"
        d.save(path)
        run("extract.py", "--input", self.src, "--output", self.out, "--workers", "1")
        corpus = json.load(open(os.path.join(self.out, "corpus.json")))
        self.assertIn("Hydrochloric acid", " ".join(corpus.values()))


class PlainTextInput(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="dsb_txt_")
        self.src = os.path.join(self.tmp, "docs")
        self.out = os.path.join(self.tmp, "out")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_markdown_is_extracted(self):
        write(os.path.join(self.src, "Math", "notes.md"),
              "# Polynomials\n\nA polynomial in one variable is an expression of a form.\n\n"
              "## Degree\n\nThe degree is the highest power of the variable.\n")
        run("extract.py", "--input", self.src, "--output", self.out, "--workers", "1")
        corpus = json.load(open(os.path.join(self.out, "corpus.json")))
        self.assertIn("polynomial", " ".join(corpus.values()).lower())

    def test_unreadable_file_is_skipped_not_fatal(self):
        os.makedirs(os.path.join(self.src, "Math"), exist_ok=True)
        with open(os.path.join(self.src, "Math", "broken.pdf"), "wb") as f:
            f.write(b"this is not a pdf at all")
        write(os.path.join(self.src, "Math", "good.txt"),
              "A circle is a closed curve whose every point is at a fixed distance.\n")
        run("extract.py", "--input", self.src, "--output", self.out, "--workers", "1")
        corpus = json.load(open(os.path.join(self.out, "corpus.json")))
        self.assertTrue(any("circle" in v.lower() for v in corpus.values()),
                        "a corrupt file broke the whole extraction")


class ConfigurableBehaviour(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="dsb_cfg_")
        self.out = os.path.join(self.tmp, "out")
        os.makedirs(self.out, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_custom_subject_keywords_are_honoured(self):
        sys.path.insert(0, SCRIPTS)
        from common import detect_subject
        cfg = {"subjects": ["Legal"], "subject_keywords": {
            "Legal": ["statute", "plaintiff", "defendant", "tribunal", "injunction"]}}
        cfgp = os.path.join(self.tmp, "c.json")
        write(cfgp, json.dumps(cfg))
        from common import load_config
        loaded = load_config(cfgp)
        self.assertIn("subject_keywords", loaded)
        text = "The plaintiff filed a case before the tribunal seeking an injunction."
        self.assertEqual(detect_subject(text, ["Legal"]), "General",
                         "built-in keywords must not apply to a custom label set")
        # with custom keywords wired in, detection should work
        from common import SUBJECT_KEYWORDS
        SUBJECT_KEYWORDS["Legal"] = cfg["subject_keywords"]["Legal"]
        try:
            self.assertEqual(detect_subject(text, ["Legal"]), "Legal")
        finally:
            SUBJECT_KEYWORDS.pop("Legal", None)

    def test_invalid_config_is_rejected_clearly(self):
        from common import load_config
        p = os.path.join(self.tmp, "bad.json")
        write(p, json.dumps({"split": {"eval_frac": 5}}))
        with self.assertRaises(SystemExit) as cm:
            load_config(p)
        self.assertIn("eval_frac", str(cm.exception))

    def test_malformed_config_json_is_rejected_clearly(self):
        from common import load_config
        p = os.path.join(self.tmp, "broken.json")
        write(p, "{not json")
        with self.assertRaises(SystemExit) as cm:
            load_config(p)
        self.assertIn("not valid JSON", str(cm.exception))

    def test_extract_limit(self):
        src = os.path.join(self.tmp, "docs")
        for i in range(6):
            write(os.path.join(src, "Math", f"f{i}.txt"),
                  f"Concept {i} is defined in this document with a full sentence here.\n")
        run("extract.py", "--input", src, "--output", self.out, "--workers", "1", "--limit", "2")
        corpus = json.load(open(os.path.join(self.out, "corpus.json")))
        self.assertLessEqual(len(corpus), 2, "--limit was ignored")


class NearDuplicateQuestions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="dsb_dupq_")
        self.out = os.path.join(self.tmp, "out")
        os.makedirs(self.out, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_validate_flags_near_duplicate_questions(self):
        base = "Define the term polynomial and explain its degree in class ten"
        rows = [
            {"text": f"<|im_start|>user\n{base}\n<|im_end|>\n<|im_start|>assistant\n"
                     f"A polynomial is an expression of a form with a degree.\n<|im_end|>",
             "subject": "Math", "instruction": base, "input": "",
             "output": "A polynomial is an expression of a form with a degree."},
            {"text": "<|im_start|>user\nDefine the term polynomial and explain its "
                     "degree in class 10\n<|im_end|>\n<|im_start|>assistant\n"
                     "A polynomial is an expression of a form with a degree.\n<|im_end|>",
             "subject": "Math",
             "instruction": "Define the term polynomial and explain its degree in class 10",
             "input": "",
             "output": "A polynomial is an expression of a form with a degree."},
        ]
        write(os.path.join(self.out, "train.jsonl"),
              "".join(json.dumps(r) + "\n" for r in rows))
        write(os.path.join(self.out, "eval.jsonl"), "")
        p = run("validate.py", "--output", self.out, expect_ok=False)
        self.assertIn("near-duplicate", p.stdout.lower())


if __name__ == "__main__":
    unittest.main(verbosity=2)
